"""把 openpyxl 读到的每个单元格按 str(v) 打印出来，作为 Go 读取器的基准。
用法：python xlsx_dump.py <file.xlsx> [max_rows]
"""
import sys
from openpyxl import load_workbook

path = sys.argv[1]
max_rows = int(sys.argv[2]) if len(sys.argv) > 2 else 10**9
wb = load_workbook(path, data_only=True)
ws = wb.worksheets[0]
for ri, row in enumerate(ws.iter_rows(values_only=True)):
    if ri >= max_rows:
        break
    for ci, c in enumerate(row):
        print(f"{ri},{ci}\t{type(c).__name__}\t{str(c)}")
