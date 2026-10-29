"""广告利润：在 ERP 财务报表口径上叠加挂载服务（jdfenxi）的 CPS / 京准通花费。

数据来源
--------
财务数据  ERP 本仓（或全仓合计）的报表口径：直接复用 report._summary_of / _summary_all，
          与「财务报表」页完全同一套计算与缓存（同一个 _heavy_cached key）。
广告花费  挂载服务 jdfenxi（默认 http://127.0.0.1:18011，frp 隧道只绑本机）的「花费一览」本地缓存：
            /api/paycomm/cache       CPS 花费（付款佣金）
            /api/paycomm/spend/cache 京准通花费
          只读缓存，**不触发**京东抓取：没有缓存的区间由前端提示去 jdfenxi 抓一次。

口径
----
    广告费合计 = CPS 花费 + 京准通花费
    广告后毛利 = 毛利 − 广告费合计
    广告后净利 = 净利润 − 广告费合计

广告费按店铺账号汇总。缓存是「一次抓取 = 一个账号 + 一个区间」的落地文件，
本接口只累加「能正好拼满所选区间」的缓存段（与 jdfenxi 月度利润页的读缓存规则一致），
拼不满的日期单列为缺口，绝不用区间外的数据凑数。
"""
from __future__ import annotations

import io
import json
import os
import re
import time
import urllib.parse
import urllib.request
from datetime import date, datetime, timedelta

from fastapi import APIRouter, Depends, HTTPException
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..auth import get_current_user
from ..database import get_db_wh, get_sessionmaker, load_warehouses, resolve_key
from ..models import User, WarehouseProduct
from .report import ALL_WH, _heavy_cached, _summary_all, _summary_of

router = APIRouter(prefix="/api", tags=["adprofit"])

# 挂载服务地址：只给后端到后端用（前端不接触这个内网端口）
AD_SERVICE_URL = os.getenv("ERP_AD_SERVICE_URL", "http://127.0.0.1:18011").rstrip("/")
AD_TIMEOUT = float(os.getenv("ERP_AD_TIMEOUT", "8"))

# 花费缓存清单变动不频繁：60 秒内复用，避免页面反复刷新时每次都打挂载服务
_AD_TTL = 60.0
_AD_CACHE: dict[str, tuple[float, object]] = {}


def _ad_get(path: str):
    """读挂载服务的 JSON 接口；异常由调用方降级处理（不能拖垮 ERP 报表）。"""
    req = urllib.request.Request(
        AD_SERVICE_URL + path,
        headers={"Accept": "application/json", "User-Agent": "WSFC-ERP"},
    )
    with urllib.request.urlopen(req, timeout=AD_TIMEOUT) as resp:  # noqa: S310 - 固定内网地址
        return json.loads(resp.read().decode("utf-8"))


def _ad_cached(key: str, loader, refresh: bool = False):
    now = time.time()
    hit = _AD_CACHE.get(key)
    if not refresh and hit and now - hit[0] < _AD_TTL:
        return hit[1]
    got = loader()
    _AD_CACHE[key] = (now, got)
    return got


def _load_entries(kind: str, refresh: bool = False) -> list[dict]:
    """花费缓存清单（按账号+区间去重）。kind: cps / spend。"""
    path = "/api/paycomm/cache" if kind == "cps" else "/api/paycomm/spend/cache"

    def _load():
        data = _ad_get(path)
        entries = data.get("entries") if isinstance(data, dict) else None
        return entries if isinstance(entries, list) else []

    return _ad_cached("entries:" + kind, _load, refresh)


def _load_accounts(refresh: bool = False) -> list[dict]:
    """挂载服务里配置的店铺账号（用于提示「有账号还没有花费缓存」）。"""
    def _load():
        data = _ad_get("/api/accounts")
        accounts = data.get("accounts") if isinstance(data, dict) else None
        return accounts if isinstance(accounts, list) else []

    return _ad_cached("accounts", _load, refresh)


def _day(value) -> date | None:
    try:
        return date.fromisoformat(str(value or "").strip()[:10])
    except ValueError:
        return None


def _span_text(s: date, e: date) -> str:
    return s.isoformat() if s == e else f"{s.isoformat()}~{e.isoformat()}"


def _tile_spans(spans: list[tuple[date, date]], start: date, end: date):
    """用缓存段正好拼满 [start, end]。

    只接受「从缺口后的第一天开始」的缓存段（段的总花费含整段金额，跨出所选区间就会多算），
    拼不满的部分返回缺口，交给调用方在前端提示。
    """
    chosen: list[tuple[date, date]] = []
    gaps: list[tuple[date, date]] = []
    ordered = sorted(spans)
    cur = start
    while cur <= end:
        best = None
        for s, e in ordered:
            if s == cur and e <= end and (best is None or e > best[1]):
                best = (s, e)
        if best:
            chosen.append(best)
            cur = best[1] + timedelta(days=1)
            continue
        nxt = min((s for s, _ in ordered if cur < s <= end), default=None)
        if nxt is None:
            gaps.append((cur, end))
            break
        gaps.append((cur, nxt - timedelta(days=1)))
        cur = nxt
    return chosen, gaps


def _spend_part(kind: str, start: date, end: date, refresh: bool):
    """某一类花费（cps / 京准通）在区间内的账号明细与合计。"""
    label = "CPS 花费" if kind == "cps" else "京准通花费"
    try:
        entries = _load_entries(kind, refresh)
    except Exception as exc:  # noqa: BLE001 - 挂载服务不可用时降级，不报 500
        print(f"[广告利润] 读取{label}缓存失败：{exc}")
        return {"ok": False, "error": f"挂载服务的「{label}」缓存读取失败，广告费未计入",
                "total": 0.0, "accounts": []}

    # 同账号同区间可能有多份历史文件：取最新抓取的那份
    by_acc: dict[str, dict[tuple[date, date], dict]] = {}
    for e in entries:
        s, t = _day(e.get("start")), _day(e.get("end"))
        if not s or not t or s > t:
            continue
        key = str(e.get("account") or "")
        slot = by_acc.setdefault(key, {})
        old = slot.get((s, t))
        if old is None or str(e.get("savedAt") or "") >= str(old.get("savedAt") or ""):
            slot[(s, t)] = e

    accounts: list[dict] = []
    total = 0.0
    for key, slots in by_acc.items():
        chosen, gaps = _tile_spans(sorted(slots), start, end)
        if not chosen:
            continue
        picked = [slots[span] for span in chosen]
        amount = round(sum(float(x.get("total") or 0) for x in picked), 2)
        total += amount
        accounts.append({
            "key": key,
            "label": str((picked[-1].get("accountLabel") or key)),
            "total": amount,
            "covered_days": sum((e - s).days + 1 for s, e in chosen),
            "spans": [_span_text(s, e) for s, e in chosen],
            "gaps": [_span_text(s, e) for s, e in gaps],
            "saved_at": max((str(x.get("savedAt") or "") for x in picked), default=""),
            "sku_rows": sum(int(x.get("rows") or 0) for x in picked),
        })
    accounts.sort(key=lambda a: -a["total"])
    return {"ok": True, "error": None, "total": round(total, 2), "accounts": accounts}


# SKU 明细要用挂载服务「花费一览」的导出接口（现成的按 SKU 报表），
# 落地文件不会变，解析结果缓存久一点（10 分钟），避免翻页/改排序都重复解析 Excel。
_SKU_TTL = 600.0
_SKU_ROWS_CACHE: dict[tuple, tuple[float, list[dict]]] = {}

# 两类花费在导出 Excel 里的金额列（列名与挂载服务导出模板一致）
_SKU_AMOUNT_COL = {"cps": "付款佣金金额(元)", "spend": "京准通花费合计(元)"}


def _ad_get_bytes(path: str) -> bytes:
    """下载挂载服务的二进制接口（导出 Excel）；超时放宽一些，导出要现装配工作簿。"""
    req = urllib.request.Request(
        AD_SERVICE_URL + path,
        headers={"Accept": "*/*", "User-Agent": "WSFC-ERP"},
    )
    with urllib.request.urlopen(req, timeout=max(AD_TIMEOUT * 2, 15)) as resp:  # noqa: S310 - 固定内网地址
        return resp.read()


def _parse_sku_xlsx(data: bytes, kind: str) -> list[dict]:
    """解析挂载服务导出的按 SKU 明细（第一个工作表：SKU / 商品名称 / 金额）。

    少于两行（只有表头）或模板变动找不到列时返回空列表，由调用方按「无数据」处理，
    不影响店铺合计口径（合计来自缓存清单，不依赖这个明细）。
    """
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        ws = wb.worksheets[0]
        header = [str(c.value or "").strip() for c in next(ws.iter_rows(min_row=1, max_row=1))]
        try:
            i_sku, i_name = header.index("SKU"), header.index("商品名称")
            i_amt = header.index(_SKU_AMOUNT_COL[kind])
        except ValueError:
            return []
        rows: list[dict] = []
        for r in ws.iter_rows(min_row=2, values_only=True):
            if not r:
                continue
            sku = str(r[i_sku] or "").strip()
            if not sku or sku == "合计":
                break
            try:
                amount = float(r[i_amt] or 0)
            except (TypeError, ValueError):
                amount = 0.0
            rows.append({
                "sku": sku,
                "name": str(r[i_name] or "").strip(),
                "amount": round(amount, 2),
            })
        return rows
    finally:
        wb.close()


def _sku_rows_of(kind: str, account: str, s: date, e: date, refresh: bool = False) -> list[dict]:
    """某账号某缓存段的 SKU 明细（按落地文件解析，失败抛异常由调用方降级）。"""
    key = ("skurows", kind, account, s.isoformat(), e.isoformat())
    now = time.time()
    hit = _SKU_ROWS_CACHE.get(key)
    if not refresh and hit and now - hit[0] < _SKU_TTL:
        return hit[1]
    qs = urllib.parse.urlencode({
        "account": account, "start": s.isoformat(), "end": e.isoformat(), "dim": "sku",
    })
    raw = _ad_get_bytes(f"/api/paycomm{'' if kind == 'cps' else '/spend'}/export?" + qs)
    rows = _parse_sku_xlsx(raw, kind)
    _SKU_ROWS_CACHE[key] = (now, rows)
    return rows


def _sku_ads_map(start: date, end: date, refresh: bool) -> tuple[dict[str, dict], list[str]]:
    """区间内各 SKU 的两类广告花费（按缓存段整段读取，拼不满的段不取）。

    返回 ({sku: {cps, jzt, name}}, warnings)；某项读取失败只降级该项，不阻塞另一项。
    """
    ads: dict[str, dict] = {}
    warnings: list[str] = []
    for kind, label in (("cps", "CPS 花费"), ("spend", "京准通花费")):
        field = "cps" if kind == "cps" else "jzt"
        try:
            entries = _load_entries(kind, refresh)
        except Exception as exc:  # noqa: BLE001 - 挂载服务不可用时降级
            print(f"[广告利润] 读取{label}缓存清单失败：{exc}")
            warnings.append(f"挂载服务的「{label}」缓存读取失败，该部分未计入 SKU 明细")
            continue

        by_acc: dict[str, dict[tuple[date, date], dict]] = {}
        for e in entries:
            s, t = _day(e.get("start")), _day(e.get("end"))
            if not s or not t or s > t:
                continue
            acc = str(e.get("account") or "")
            slot = by_acc.setdefault(acc, {})
            old = slot.get((s, t))
            if old is None or str(e.get("savedAt") or "") >= str(old.get("savedAt") or ""):
                slot[(s, t)] = e

        for acc, slots in by_acc.items():
            chosen, _gaps = _tile_spans(sorted(slots), start, end)
            for span in chosen:
                try:
                    rows = _sku_rows_of(kind, acc, span[0], span[1], refresh)
                except Exception as exc:  # noqa: BLE001
                    print(f"[广告利润] 读取 {acc} 的{label} SKU 明细失败：{exc}")
                    warnings.append(f"{label}（{acc}）{_span_text(*span)} 的 SKU 明细读取失败，已跳过")
                    continue
                for r in rows:
                    d = ads.setdefault(r["sku"], {"sku": r["sku"], "ad_name": "", "cps": 0.0, "jzt": 0.0})
                    d[field] = round(d[field] + r["amount"], 2)
                    if r["name"] and not d["ad_name"]:
                        d["ad_name"] = r["name"]
    return ads, warnings


# ---------------------------------------------------------------- 月度利润（入仓品）
# 入仓品的成交件数/毛利只有 jdfenxi「月度利润」里有（商智成交 + ERP 成本接口算出的毛利），
# 这里只读它的本地缓存（export 接口，缓存缺失时返回最新一份，用文件名里的区间校验），
# 不触发商智抓取；同一个月只有部分日期有缓存时按已有区间取数并在页面上标注截止日。
_PROFIT_TTL = 600.0
_PROFIT_CACHE: dict[tuple, tuple[float, dict]] = {}

# 月度利润 Excel「SKU利润」表的列名
_PROFIT_COLS = {
    "sku": "SKU",
    "name": "商品名称",
    "qty": "成交件数",
    "gross": "毛利(元)",
    "arrival": "到手金额(元)",
}


def _ad_get_with_headers(path: str) -> tuple[bytes, dict]:
    req = urllib.request.Request(
        AD_SERVICE_URL + path,
        headers={"Accept": "*/*", "User-Agent": "WSFC-ERP"},
    )
    with urllib.request.urlopen(req, timeout=max(AD_TIMEOUT * 2, 20)) as resp:  # noqa: S310
        # Starlette 下发的响应头是小写，统一转小写再取（别按 "Content-Disposition" 取）
        return resp.read(), {k.lower(): v for k, v in resp.headers.items()}


def _xlsx_response_range(disposition: str) -> tuple[date, date] | None:
    """从导出文件名里取真实数据区间（月度利润_账号_开始_结束.xlsx）。"""
    if not disposition:
        return None
    m = re.search(r"filename\*=UTF-8''([^;]+)", disposition)
    name = urllib.parse.unquote(m.group(1)) if m else disposition
    m2 = re.search(r"(\d{4}-\d{2}-\d{2})_(\d{4}-\d{2}-\d{2})\.xlsx", name)
    if not m2:
        return None
    s, e = _day(m2.group(1)), _day(m2.group(2))
    return (s, e) if s and e else None


def _parse_profit_xlsx(data: bytes) -> list[dict]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        ws = wb.worksheets[0]
        header = [str(c.value or "").strip() for c in next(ws.iter_rows(min_row=1, max_row=1))]
        idx = {}
        for k, label in _PROFIT_COLS.items():
            try:
                idx[k] = header.index(label)
            except ValueError:
                return []
        rows: list[dict] = []
        for r in ws.iter_rows(min_row=2, values_only=True):
            if not r:
                continue
            sku = str(r[idx["sku"]] or "").strip()
            if not sku or sku == "合计":
                break
            try:
                qty = float(r[idx["qty"]] or 0)
                gross = float(r[idx["gross"]] or 0)
                arrival = float(r[idx["arrival"]] or 0)
            except (TypeError, ValueError):
                continue
            rows.append({
                "sku": sku,
                "name": str(r[idx["name"]] or "").strip(),
                "qty": qty,
                "gross": round(gross, 2),
                "arrival": round(arrival, 2),
            })
        return rows
    finally:
        wb.close()


def _profit_month(month_start: date, month_end: date, refresh: bool) -> tuple[dict[str, dict], list[str], dict]:
    """某个月各 SKU 的成交件数/毛利（各账号取本地利润缓存，同 SKU 取成交件数大的一份）。"""
    key = ("profit", month_start.isoformat())
    now = time.time()
    hit = _PROFIT_CACHE.get(key)
    if not refresh and hit and now - hit[0] < _PROFIT_TTL:
        return hit[1]["rows"], hit[1]["warnings"], hit[1]["meta"]

    accounts = ["main", "b"]
    rows: dict[str, dict] = {}
    warnings: list[str] = []
    used_ranges: list[str] = []
    rng: tuple[date, date] | None = None
    for acc in accounts:
        # 第一次先问 main，从导出文件名拿到该月真正缓存到的区间；
        # 之后各账号用这个精确区间去读，才能命中缓存并按账号拆行（否则 export 会回落成 main 的表）
        req_s, req_e = (month_start, month_end) if rng is None else rng
        qs = urllib.parse.urlencode({
            "account": acc, "start": req_s.isoformat(),
            "end": req_e.isoformat(), "dim": "sku",
        })
        try:
            raw, headers = _ad_get_with_headers("/api/profit/export?" + qs)
        except Exception as exc:  # noqa: BLE001 - 挂载服务不可用/无缓存
            print(f"[广告利润] 读取 {acc} 月度利润缓存失败：{exc}")
            continue
        got = _xlsx_response_range(headers.get("content-disposition", ""))
        if not got:
            continue
        if rng is None:
            # export 找不到该月缓存时会回落「最新一份」：只接受落在本月内的区间，避免拿错月份的数
            if got[0] < month_start or got[1] > month_end:
                continue
            rng = got
        elif got != rng:
            continue          # 读到的不是同一份缓存（该账号没有这张表）
        used_ranges.append(f"{got[0].isoformat()}~{got[1].isoformat()}")
        try:
            part = _parse_profit_xlsx(raw)
        except Exception as exc:  # noqa: BLE001
            print(f"[广告利润] 解析 {acc} 月度利润失败：{exc}")
            continue
        for r in part:
            cur = rows.get(r["sku"])
            # 同一 SKU 两个账号都有成交时不要相加（会重复计），取成交件数大的那份
            if cur is None or r["qty"] > cur["qty"]:
                rows[r["sku"]] = r

    if not used_ranges:
        warnings.append(
            f"{month_start.strftime('%Y-%m')} 没有月度利润缓存，入仓品的成交件数/毛利未补全"
            "（到 jdfenxi「月度利润」页算一次即可）"
        )
    else:
        r_s = min(x.split("~")[0] for x in used_ranges)
        r_e = max(x.split("~")[1] for x in used_ranges)
        if r_s != month_start.isoformat() or r_e != month_end.isoformat():
            warnings.append(
                f"入仓品的成交件数/毛利取自月度利润缓存 {r_s}~{r_e}"
                f"（{month_start.strftime('%Y-%m')} 未覆盖到最后一天）"
            )

    meta = {"label": "、".join(sorted(set(used_ranges)))}
    out = (rows, warnings, meta)
    _PROFIT_CACHE[key] = (now, {"rows": rows, "warnings": warnings, "meta": meta})
    return out


def _profit_for_range(start: date, end: date, refresh: bool) -> tuple[dict[str, dict], list[str], list[str]]:
    """所选区间覆盖的各月利润数据（按 SKU 合并各月）；返回 (rows, warnings, 用到的缓存区间)。"""
    rows: dict[str, dict] = {}
    warnings: list[str] = []
    ranges: list[str] = []
    cur = start.replace(day=1)
    while cur <= end:
        nxt = (cur.replace(day=1) + timedelta(days=32)).replace(day=1)
        m_end = nxt - timedelta(days=1)
        m_rows, m_warn, meta = _profit_month(cur, m_end, refresh)
        for sku, r in m_rows.items():
            cur_row = rows.get(sku)
            if cur_row is None:
                rows[sku] = dict(r)
            else:
                cur_row["qty"] += r["qty"]
                cur_row["gross"] = round(cur_row["gross"] + r["gross"], 2)
                cur_row["arrival"] = round(cur_row["arrival"] + r["arrival"], 2)
        warnings.extend(m_warn)
        if meta.get("label"):
            ranges.extend(meta["label"].split("、"))
        cur = nxt
    return rows, warnings, sorted(set(ranges))


def _warehouse_skus(wh_key: str, db: Session) -> dict[str, str]:
    """当前口径下的入仓品资料：SKU → 名称（全仓时把所有分仓的资料并起来）。"""

    def _scan(session) -> dict[str, str]:
        out: dict[str, str] = {}
        for p in session.execute(select(WarehouseProduct)).scalars():
            sku = str(p.sku or "").strip()
            if sku:
                out.setdefault(sku, p.name or "")
        return out

    if wh_key != ALL_WH:
        return _scan(db)
    out: dict[str, str] = {}
    for w in load_warehouses().get("list") or []:
        try:
            session = get_sessionmaker(str(w.get("key")) )()
        except Exception:  # noqa: BLE001 - 某个仓文件坏了不影响其他仓
            continue
        try:
            for sku, name in _scan(session).items():
                out.setdefault(sku, name)
        finally:
            session.close()
    return out


def _parse_range(date_from: str, date_to: str) -> tuple[date, date]:
    """接口入参的日期区间（默认本月到今天，最长 366 天）。"""
    today = date.today()
    start = _day(date_from) if date_from else today.replace(day=1)
    end = _day(date_to) if date_to else today
    if not start or not end:
        raise HTTPException(400, "日期格式不对（要 YYYY-MM-DD）")
    if start > end:
        start, end = end, start
    if (end - start).days > 366:
        raise HTTPException(400, "区间最长 366 天")
    return start, end


@router.get("/adprofit/summary")
def ad_profit_summary(
    date_from: str = "",
    date_to: str = "",
    wh: str = "",                 # 空 = 本登录会话分仓；wh=all = 全仓合计
    exclude_other: int = 1,       # 与「财务报表」默认口径一致：排除其他收支
    refresh: int = 0,             # 1 = 忽略 60 秒缓存，重新读挂载服务
    db: Session = Depends(get_db_wh),
    user: User = Depends(get_current_user),
):
    """含广告费的利润：ERP 财务口径 −（CPS 花费 + 京准通花费）。"""
    start, end = _parse_range(date_from, date_to)

    s_iso, e_iso = start.isoformat(), end.isoformat()
    wh_key = ALL_WH if wh == ALL_WH else resolve_key(wh)
    key = ("summary", wh_key, s_iso, e_iso, 1 if exclude_other else 0)

    def _compute():
        if wh_key == ALL_WH:
            return _summary_all(s_iso, e_iso, bool(exclude_other))
        return _summary_of(db, s_iso, e_iso, wh_key, bool(exclude_other))

    base = _heavy_cached(key, _compute)

    cps = _spend_part("cps", start, end, bool(refresh))
    jzt = _spend_part("spend", start, end, bool(refresh))
    service_ok = cps["ok"] and jzt["ok"]

    ads_total = round(cps["total"] + jzt["total"], 2)
    revenue = round(float(base.get("revenue") or 0), 2)
    gross = round(float(base.get("gross_profit") or 0), 2)
    net = round(float(base.get("net_profit") or 0), 2)
    days = (end - start).days + 1

    # ---- 覆盖提示：拼不满的日期必须在页面上说清楚，否则广告费会被低估 ----
    warnings: list[str] = []
    covered_accounts: dict[str, set[str]] = {}
    for part_name, part in (("CPS 花费", cps), ("京准通花费", jzt)):
        for acc in part["accounts"]:
            if acc["gaps"]:
                covered_accounts.setdefault(acc["label"], set()).add(
                    f"{part_name}缺 " + "、".join(acc["gaps"])
                )
    for label, items in covered_accounts.items():
        warnings.append(f"{label}：" + "；".join(sorted(items)))
    if not service_ok:
        if not cps["ok"]:
            warnings.append(cps["error"])
        if not jzt["ok"]:
            warnings.append(jzt["error"])
    elif ads_total <= 0:
        warnings.append("所选区间没有任何可用的花费缓存，广告费按 0 计入；请到挂载服务「花费一览」按该区间抓取一次")

    # ---- 有账号但该区间没有缓存：单列出来，提醒去抓取 ----
    # 只提醒「曾经抓过花费」的账号（有缓存文件、但没覆盖本区间），
    # 一次都没抓过的账号多半没在投广告，天天提示反而淹没真正需要补抓的。
    missing_accounts: list[dict] = []
    if service_ok:
        try:
            accounts = _load_accounts(bool(refresh))
        except Exception:  # noqa: BLE001
            accounts = []
        known: set[str] = set()
        for kind in ("cps", "spend"):
            try:
                for e in _load_entries(kind, bool(refresh)):
                    k = str(e.get("account") or "")
                    if k:
                        known.add(k)
            except Exception:  # noqa: BLE001
                continue
        seen = {a["key"] for a in cps["accounts"]} | {a["key"] for a in jzt["accounts"]}
        for acc in accounts:
            key_ = str(acc.get("key") or "")
            if key_ and key_ not in seen and key_ in known:
                missing_accounts.append({"key": key_, "label": str(acc.get("label") or key_)})
        if missing_accounts:
            warnings.append("以下店铺账号该区间没有花费缓存（未计入广告费）："
                            + "、".join(a["label"] for a in missing_accounts))

    result = {
        "ok": True,
        "generated_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "range": {"start": s_iso, "end": e_iso, "days": days},
        "erp": {
            "warehouse": base.get("warehouse") or {"key": wh_key, "name": ""},
            "is_all": bool(base.get("is_all")),
            "revenue": revenue,
            "cogs": round(float(base.get("cogs") or 0), 2),
            "goods_cogs": round(float(base.get("goods_cogs") or 0), 2),
            "brush_cost": round(float(base.get("brush_cost") or 0), 2),
            "gross_profit": gross,
            "expense": round(float(base.get("expense") or 0), 2),
            "manual_expense": round(float(base.get("manual_expense") or 0), 2),
            "other_expense": round(float(base.get("other_expense") or 0), 2),
            "excluded_other_expense": round(float(base.get("excluded_other_expense") or 0), 2),
            "exclude_other_expense": bool(base.get("exclude_other_expense")),
            "net_profit": net,
            "order_count": int(base.get("order_count") or 0),
            "warehouse_in_revenue": round(float(base.get("warehouse_in_revenue") or 0), 2),
            "warehouse_in_cogs": round(float(base.get("warehouse_in_cogs") or 0), 2),
            "warehouse_in_count": int(base.get("warehouse_in_count") or 0),
        },
        "ads": {
            "service_ok": service_ok,
            "total": ads_total,
            "cps": cps,
            "jzt": jzt,
            "missing_accounts": missing_accounts,
        },
        "profit": {
            "ads_total": ads_total,
            "cps_total": cps["total"],
            "jzt_total": jzt["total"],
            "gross_before_ads": gross,
            "gross_after_ads": round(gross - ads_total, 2),
            "net_before_ads": net,
            "net_after_ads": round(net - ads_total, 2),
            "ad_rate": round(ads_total / revenue, 4) if revenue else None,
            "net_margin_before": round(net / revenue, 4) if revenue else None,
            "net_margin_after": round((net - ads_total) / revenue, 4) if revenue else None,
        },
        "warnings": warnings,
    }
    return result


@router.get("/adprofit/skus")
def ad_profit_skus(
    date_from: str = "",
    date_to: str = "",
    wh: str = "",
    exclude_other: int = 1,
    refresh: int = 0,
    sort: str = "ads",            # ads=按广告费倒序（默认） / revenue / gross
    db: Session = Depends(get_db_wh),
    user: User = Depends(get_current_user),
):
    """按 SKU 的含广告费利润：ERP 商品明细（按商品编码=京东 skuId 合并）+ 挂载服务按 SKU 的花费。

    匹配口径：ERP 商品的「编码」= 挂载服务花费明细里的 SKU（京东 skuId）。
    没填编码的商品按商品名称单列（不参与广告费匹配）；只有广告花费、没有 ERP 销售的 SKU
    也单独列出（收入为 0），方便看出「广告在烧、货没卖」的款。
    """
    start, end = _parse_range(date_from, date_to)
    s_iso, e_iso = start.isoformat(), end.isoformat()
    wh_key = ALL_WH if wh == ALL_WH else resolve_key(wh)
    key = ("summary", wh_key, s_iso, e_iso, 1 if exclude_other else 0)

    def _compute():
        if wh_key == ALL_WH:
            return _summary_all(s_iso, e_iso, bool(exclude_other))
        return _summary_of(db, s_iso, e_iso, wh_key, bool(exclude_other))

    # 与 summary 共用同一个缓存键：SKU 页不额外算一遍财务数据
    base = _heavy_cached(key, _compute)
    ads, warnings = _sku_ads_map(start, end, bool(refresh))

    # 入仓品：ERP 里没有它们的出库销售（只按袋入仓），成交件数/毛利只有 jdfenxi 月度利润有
    wh_skus = _warehouse_skus(wh_key, db)
    profit_rows, profit_warn, profit_ranges = _profit_for_range(start, end, bool(refresh))
    warnings.extend(profit_warn)

    # ---- ERP 商品明细按「编码」合并（同编码多规格/代发/入仓合并为一行）----
    buckets: dict[tuple[str, str], dict] = {}
    for p in base.get("by_product") or []:
        # 入仓台账（备货入库）不是 SKU 成交，且没有商品编码；入仓品由下面的月度利润行按 SKU 展示，
        # 两边都算会把同一批货的「入仓采购结算」和「京东成交」重复加进合计
        if p.get("is_warehouse_in"):
            continue
        code = str(p.get("code") or "").strip()
        bkey = ("sku", code) if code else ("name", str(p.get("name") or ""))
        b = buckets.get(bkey)
        if b is None:
            b = buckets[bkey] = {
                "code": code, "name": str(p.get("name") or ""), "unit": "",
                "qty": 0.0, "revenue": 0.0, "cogs": 0.0, "gross_before": 0.0,
            }
        b["qty"] += float(p.get("qty") or 0)
        # 收入口径与报表一致：商品销售金额 + 客户代收的关联结算（包材/人工/快递）
        b["revenue"] += float(p.get("amount") or 0) + float(p.get("settle_income") or 0)
        b["cogs"] += float(p.get("total_cogs") or 0)
        b["gross_before"] += float(p.get("gross_profit") or 0)
        if not b["unit"] and p.get("unit"):
            b["unit"] = str(p["unit"])

    rows: list[dict] = []
    matched: set[str] = set()
    for b in buckets.values():
        a = ads.get(b["code"]) if b["code"] else None
        if b["code"] and a:
            matched.add(b["code"])
        cps = round(float((a or {}).get("cps") or 0), 2)
        jzt = round(float((a or {}).get("jzt") or 0), 2)
        ads_total = round(cps + jzt, 2)
        revenue = round(b["revenue"], 2)
        gross = round(b["gross_before"], 2)
        rows.append({
            "sku": b["code"],
            "name": b["name"],
            "qty": round(b["qty"], 4),
            "unit": b["unit"],
            "revenue": revenue,
            "cogs": round(b["cogs"], 2),
            "gross_before": gross,
            "cps": cps,
            "jzt": jzt,
            "ads_total": ads_total,
            "gross_after": round(gross - ads_total, 2),
            "ad_rate": round(ads_total / revenue, 4) if revenue else None,
            "roi": round(revenue / ads_total, 4) if ads_total else None,
            "has_sales": False if abs(revenue) < 0.005 else True,
            "no_code": not b["code"],
            "ad_only": False,
            "is_warehouse_in": False,
            "source": "erp",
        })

    # ---- 入仓品：成交件数/毛利取自月度利润缓存，利润 = 入仓品毛利 − 广告费 ----
    erp_codes = {b["code"] for b in buckets.values() if b["code"]}
    for sku, pr in profit_rows.items():
        if sku in erp_codes or sku not in wh_skus:
            continue          # 已在 ERP 销售明细里 / 不是入仓品
        a = ads.get(sku)
        if a:
            matched.add(sku)
        cps = round(float((a or {}).get("cps") or 0), 2)
        jzt = round(float((a or {}).get("jzt") or 0), 2)
        ads_total = round(cps + jzt, 2)
        revenue = round(float(pr["arrival"] or 0), 2)
        gross = round(float(pr["gross"] or 0), 2)
        rows.append({
            "sku": sku,
            "name": wh_skus.get(sku) or pr.get("name") or "",
            "qty": round(float(pr["qty"] or 0), 4),
            "unit": "件",      # 商智成交件数，不是入仓袋数
            "revenue": revenue,
            "cogs": round(revenue - gross, 2),
            "gross_before": gross,
            "cps": cps,
            "jzt": jzt,
            "ads_total": ads_total,
            "gross_after": round(gross - ads_total, 2),
            "ad_rate": round(ads_total / revenue, 4) if revenue else None,
            "roi": round(revenue / ads_total, 4) if ads_total else None,
            "has_sales": True,
            "no_code": False,
            "ad_only": False,
            "is_warehouse_in": True,
            "source": "profit",
        })

    # ---- 只有广告花费、没有 ERP 销售的 SKU（收入 0，单独列出）----
    for sku, a in ads.items():
        if sku in matched:
            continue
        cps = round(float(a.get("cps") or 0), 2)
        jzt = round(float(a.get("jzt") or 0), 2)
        ads_total = round(cps + jzt, 2)
        if not ads_total:
            continue
        rows.append({
            "sku": sku, "name": wh_skus.get(sku) or a.get("ad_name") or "",
            "qty": 0.0, "unit": "",
            "revenue": 0.0, "cogs": 0.0, "gross_before": 0.0,
            "cps": cps, "jzt": jzt, "ads_total": ads_total,
            "gross_after": round(-ads_total, 2),
            "ad_rate": None, "roi": None,
            "has_sales": False, "no_code": False, "ad_only": True,
            "is_warehouse_in": sku in wh_skus,
            "source": "ad",
        })

    sorters = {
        "revenue": lambda r: (-r["revenue"], -r["ads_total"]),
        "gross": lambda r: (-r["gross_before"], -r["ads_total"]),
    }
    rows.sort(key=sorters.get(sort, lambda r: (-r["ads_total"], -r["revenue"])))

    erp_revenue = round(sum(r["revenue"] for r in rows if r["source"] == "erp"), 2)
    wh_rows = [r for r in rows if r["is_warehouse_in"]]
    ads_total = round(sum(float(a.get("cps") or 0) + float(a.get("jzt") or 0) for a in ads.values()), 2)
    matched_ads = round(sum(float(ads[k].get("cps") or 0) + float(ads[k].get("jzt") or 0) for k in matched), 2)
    unmatched = round(ads_total - matched_ads, 2)
    if unmatched > 0.005:
        warnings.append(
            f"有 {fmt_money(unmatched)} 广告花费的 SKU 在 ERP 商品明细里找不到对应编码（京东 skuId），"
            "这部分已单列为「未匹配 ERP」行；如商品编码没填京东 skuId，请到商品资料补上"
        )

    return {
        "ok": True,
        "generated_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "range": {"start": s_iso, "end": e_iso, "days": (end - start).days + 1},
        "warehouse": base.get("warehouse") or {"key": wh_key, "name": ""},
        "is_all": bool(base.get("is_all")),
        "profit_ranges": profit_ranges,
        "rows": rows,
        "totals": {
            "sku_count": len(rows),
            "erp_revenue": erp_revenue,
            "revenue_total": round(sum(r["revenue"] for r in rows), 2),
            "warehouse_sku_count": len(wh_rows),
            "warehouse_qty": round(sum(r["qty"] for r in wh_rows), 4),
            "warehouse_revenue": round(sum(r["revenue"] for r in wh_rows), 2),
            "warehouse_gross": round(sum(r["gross_before"] for r in wh_rows), 2),
            "warehouse_ads": round(sum(r["ads_total"] for r in wh_rows), 2),
            "warehouse_net": round(sum(r["gross_after"] for r in wh_rows), 2),
            "gross_before": round(sum(r["gross_before"] for r in rows), 2),
            "cps_total": round(sum(r["cps"] for r in rows), 2),
            "jzt_total": round(sum(r["jzt"] for r in rows), 2),
            "ads_total": round(sum(r["ads_total"] for r in rows), 2),
            "gross_after": round(sum(r["gross_after"] for r in rows), 2),
            "matched_skus": len(matched),
            "ad_only_skus": sum(1 for r in rows if r["ad_only"]),
            "unmatched_ads": unmatched,
        },
        "warnings": warnings,
    }


def fmt_money(v: float) -> str:
    return f"¥{float(v or 0):,.2f}"
